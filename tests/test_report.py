from docx import Document

from report import build_report as br


def _text(path):
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for r in t.rows for c in r.cells]
    return "\n".join(parts)


def test_full_report_uses_computed_values(results, tmp_path):
    out = tmp_path / "out"
    assert br.main(["--results", str(results[0]), "--out", str(out), "--no-pdf"]) == 0
    txt = _text(out / "BDA_Mini_Project_Report.docx")
    s = results[1]
    for needle in (f"{s['rows_clean_unique']:,}", "Delhi", "10. References", "Table of Contents".upper(), "Hotspots", "Largest PM2.5 spike days"):
        assert needle in txt, needle
    assert "DRAFT TEMPLATE" not in txt
    assert (out / "figures" / "cities.png").exists() and (out / "figures" / "architecture.png").exists()


def test_template_marks_results_pending(tmp_path):
    out = tmp_path / "t"
    assert br.main(["--template", "--out", str(out), "--no-pdf"]) == 0
    txt = _text(out / "BDA_Mini_Project_Report_TEMPLATE.docx")
    assert "DRAFT TEMPLATE" in txt and "[Pending" in txt and "Delhi" not in txt


def test_missing_results_returns_error(tmp_path):
    assert br.main(["--results", str(tmp_path / "none"), "--out", str(tmp_path / "o")]) == 2
