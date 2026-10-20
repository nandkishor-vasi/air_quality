"""Build the DOCX (and PDF) project report from REAL pipeline outputs.

    python report/build_report.py --results results            # full report (needs a pipeline run)
    python report/build_report.py --template                   # skeleton: result sections marked pending

Nothing is invented: every number comes from results/summary.json, results/parquet/*, docs/pytest_report.xml.
Screenshots are embedded only if you captured them to docs/screenshots/ (names listed in SHOTS).
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from report import diagrams  # noqa: E402
from src.aq import config as C  # noqa: E402
from src.aq import insights as INS  # noqa: E402

SHOTS = {  # file -> caption
    "01_overview.png": "Dashboard overview: metric cards and city/state rankings.",
    "02_trends.png": "Dashboard trends: monthly series and seasonal comparison.",
    "03_hourly.png": "Dashboard hourly (diurnal) pattern for selected cities.",
    "04_hotspots.png": "Dashboard hotspot map and ranked station table.",
    "05_anomalies.png": "Dashboard anomaly view.",
    "06_aqi.png": "Dashboard AQI categories and dominant pollutants.",
    "07_data_quality.png": "Dashboard data-quality tab.",
}
PENDING = "[Pending: generated automatically after the pipeline is run on the downloaded dataset — see README.]"


# ----------------------------------------------------------------------------- charts
def make_charts(T, S, outdir: Path, param="PM2.5") -> dict:
    outdir.mkdir(parents=True, exist_ok=True)
    out = {}
    blue, orange = "#2b7bba", "#e8833a"

    def save(fig, name):
        p = outdir / f"{name}.png"
        fig.savefig(p, dpi=170, bbox_inches="tight"); plt.close(fig); out[name] = p

    cs = T["city_summary"]; cs = cs[(cs.parameter_name == param) & (cs.n_station_days >= 30)].sort_values("mean_daily", ascending=False)
    if len(cs):
        d = cs.head(10).iloc[::-1]
        fig, ax = plt.subplots(figsize=(7, 4)); ax.barh(d.city_name, d.mean_daily, color=blue)
        if param in C.NAAQS_24H:
            ax.axvline(C.NAAQS_24H[param], color=orange, ls="--", label=f"NAAQS 24 h = {C.NAAQS_24H[param]}"); ax.legend()
        ax.set_xlabel(f"Mean daily {param} (µg/m³)"); ax.set_title(f"Ten most polluted cities by mean daily {param}"); save(fig, "cities")
    ss = T["state_summary"]; ss = ss[(ss.parameter_name == param) & (ss.n_station_days >= 30)].sort_values("mean_daily", ascending=False)
    if len(ss):
        d = ss.head(10).iloc[::-1]
        fig, ax = plt.subplots(figsize=(7, 4)); ax.barh(d.state_name, d.mean_daily, color=blue)
        ax.set_xlabel(f"Mean daily {param} (µg/m³)"); ax.set_title(f"Ten highest states by mean daily {param}"); save(fig, "states")
    mn = T["monthly_national"]; mn = mn[mn.parameter_name == param].sort_values("ym")
    if len(mn):
        fig, ax = plt.subplots(figsize=(7, 3.5)); ax.plot(mn.ym, mn.mean_daily, marker="o", color=blue)
        ax.set_ylabel("µg/m³"); ax.set_title(f"Monthly mean daily {param} across all stations (unweighted by city)")
        ax.tick_params(axis="x", rotation=60); save(fig, "monthly")
    sc = T["seasonal_city"]; sc = sc[sc.parameter_name == param]
    if len(sc):
        sea = (sc.assign(x=sc.mean_daily * sc.n_station_days).groupby("season")[["x", "n_station_days"]].sum())
        sea = (sea.x / sea.n_station_days).reindex(C.SEASON_ORDER).dropna()
        fig, ax = plt.subplots(figsize=(5.5, 3.5)); ax.bar(sea.index, sea.values, color=blue)
        ax.set_ylabel("µg/m³"); ax.set_title(f"Seasonal mean daily {param}"); save(fig, "seasonal")
    hp = T["hourly_pattern"]; hp = hp[hp.parameter_name == param]
    if len(hp):
        h = hp.assign(x=hp.mean_value * hp.n_readings).groupby("hour")[["x", "n_readings"]].sum(); h = h.x / h.n_readings
        fig, ax = plt.subplots(figsize=(7, 3.5)); ax.plot(h.index, h.values, marker="o", color=blue)
        ax.set_xlabel("Hour of day (IST)"); ax.set_ylabel("µg/m³"); ax.set_title(f"Average hourly {param} profile"); ax.set_xticks(range(0, 24, 2)); save(fig, "hourly")
    hs = T["hotspots"]; hs = hs[hs.parameter_name == param]
    if len(hs) and "latitude" in hs and hs.latitude.notna().any():
        g = hs.dropna(subset=["latitude", "longitude"])
        fig, ax = plt.subplots(figsize=(5.5, 6))
        sc_ = ax.scatter(g.longitude, g.latitude, c=g.mean_daily, s=18 + 3 * g.is_hotspot.astype(int) * 10, cmap="YlOrRd", edgecolor="#555", lw=.3)
        h = g[g.is_hotspot]; ax.scatter(h.longitude, h.latitude, facecolors="none", edgecolors="black", s=70, lw=1, label="hotspot")
        fig.colorbar(sc_, label=f"Mean daily {param}"); ax.legend(); ax.set_xlabel("Longitude"); ax.set_ylabel("Latitude")
        ax.set_title("Station locations (no basemap)"); save(fig, "hotspot_map")
    an = T["anomaly_summary"]; an = an[an.parameter_name == param].sort_values("n_anomalies", ascending=False).head(10)
    if len(an) and an.n_anomalies.sum() > 0:
        fig, ax = plt.subplots(figsize=(7, 3.5)); ax.bar(an.city_name, an.n_anomalies, color=orange)
        ax.set_ylabel("Anomalous station-days"); ax.set_title(f"Cities with most {param} spikes"); ax.tick_params(axis="x", rotation=60); save(fig, "anomalies")
    cat = T.get("aqi_category_city")
    if cat is not None and len(cat):
        t = cat.groupby("aqi_category").n_station_days.sum().reindex(["Good", "Satisfactory", "Moderate", "Poor", "Very Poor", "Severe"]).dropna()
        fig, ax = plt.subplots(figsize=(6, 3.5)); ax.bar(t.index, t.values, color=["#2e9e4f", "#8cc63f", "#f0c419", "#f08a1c", "#e0451f", "#8b1a1a"][:len(t)])
        ax.set_ylabel("Station-days"); ax.set_title("Computed CPCB AQI categories"); ax.tick_params(axis="x", rotation=30); save(fig, "aqi")
    return out


# ----------------------------------------------------------------------------- docx helpers
class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.styles["Normal"]; s.font.name = "Times New Roman"; s.font.size = Pt(11)
        s.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        for n, sz in (("Heading 1", 14), ("Heading 2", 12)):
            h = self.d.styles[n]; h.font.name = "Times New Roman"; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = RGBColor(0, 0, 0)
            rf = h.element.rPr.rFonts
            for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(att)) is not None:
                    del rf.attrib[qn(att)]                      # theme fonts would override Times New Roman
            for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(att), "Times New Roman")
        sec = self.d.sections[0]
        sec.left_margin = sec.right_margin = Cm(2.3); sec.top_margin = sec.bottom_margin = Cm(2.2)
        self.fig = self.tab = 0
        self._border(sec); self._footer(sec)

    def _border(self, sec):
        pb = OxmlElement("w:pgBorders"); pb.set(qn("w:offsetFrom"), "page")
        for side in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "double"); e.set(qn("w:sz"), "6"); e.set(qn("w:space"), "24"); e.set(qn("w:color"), "000000"); pb.append(e)
        sec._sectPr.append(pb)

    def _footer(self, sec):
        p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if t:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
            else:
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
            r._r.append(e)

    def center(self, text, size=11, bold=False, color=None, after=6):
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        return p

    def h(self, text, level):
        return self.d.add_heading(text, level)

    def p(self, text, bold_prefix=None, italic=False):
        para = self.d.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; para.paragraph_format.space_after = Pt(6)
        if bold_prefix:
            para.add_run(bold_prefix + " ").bold = True
        r = para.add_run(text); r.italic = italic
        return para

    def bullets(self, items):
        for i in items:
            self.d.add_paragraph(i, style="List Bullet")

    def code(self, text):
        for line in text.strip("\n").split("\n"):
            p = self.d.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.left_indent = Cm(0.6)
            r = p.add_run(line); r.font.name = "Courier New"; r.font.size = Pt(9)
        self.d.add_paragraph()

    def figure(self, path, caption, width=14.5):
        self.fig += 1
        self.d.add_picture(str(path), width=Cm(width)); self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = self.d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.add_run(f"Fig. {self.fig}. ").bold = True; c.add_run(caption).font.size = Pt(10)

    def table(self, caption, df: pd.DataFrame, widths=None, size=9):
        self.tab += 1
        c = self.d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.keep_with_next = True
        c.add_run(f"Table {self.tab}. ").bold = True; c.add_run(caption).font.size = Pt(10)
        t = self.d.add_table(rows=1, cols=len(df.columns)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, col in enumerate(df.columns):
            cell = t.rows[0].cells[i]; cell.text = ""; r = cell.paragraphs[0].add_run(str(col)); r.bold = True; r.font.size = Pt(size)
            sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "EEEEEE"); cell._tc.get_or_add_tcPr().append(sh)
        for row in df.itertuples(index=False):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                txt = f"{v:,.2f}".rstrip("0").rstrip(".") if isinstance(v, float) else (f"{v:,}" if isinstance(v, int) else str(v))
                cells[i].text = ""; cells[i].paragraphs[0].add_run(txt).font.size = Pt(size)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.d.add_paragraph()

    def pagebreak(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ----------------------------------------------------------------------------- content
def junit_rows(path: Path):
    if not path.exists():
        return None, None
    root = ET.parse(path).getroot()
    suite = root if root.tag == "testsuite" else root.find("testsuite")
    rows = []
    for tc in suite.iter("testcase"):
        bad = tc.find("failure") is not None or tc.find("error") is not None
        rows.append((tc.get("classname", "").replace("tests.", ""), tc.get("name"), "FAILED" if bad else "passed"))
    return rows, dict(suite.attrib)


def pick(ins, *prefixes):
    return " ".join(x for x in (ins or {}).get("findings", []) if x.startswith(prefixes))


def build(cfg, results: Path | None, toc_pages: dict, out_docx: Path, figdir: Path):
    figdir.mkdir(parents=True, exist_ok=True)
    S = json.loads((results / "summary.json").read_text()) if results else None
    T = INS.load(results) if results else None
    ins = INS.generate(results) if results else None
    ch = make_charts(T, S, figdir) if results else {}
    diagrams.architecture(figdir / "architecture.png"); diagrams.pipeline(figdir / "pipeline.png")
    jr, jattr = junit_rows(ROOT / "docs" / "pytest_report.xml")
    D = Doc()
    heads: list[tuple[int, str]] = []

    def H1(t): heads.append((1, t)); D.h(t, 1)
    def H2(t): heads.append((2, t)); D.h(t, 2)
    def pend(): D.p(PENDING, italic=True)

    # cover
    D.center(cfg["trust"], 14, True); D.center(cfg["college"], 16, True); D.center(cfg["affiliation"], 10, True)
    D.center(cfg["department"], 11, True, "4F81BD", after=18)
    if cfg.get("logo_path") and Path(cfg["logo_path"]).exists():
        D.d.add_picture(cfg["logo_path"], width=Cm(4)); D.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        D.center("[Insert college logo here — set logo_path in report_config.json]", 9, after=36)
    D.center("MINI PROJECT REPORT", 12, True); D.center("on", 11, True)
    D.center(cfg["title"], 13, True, "4F81BD", after=14); D.center("Submitted by", 11)
    st = D.d.add_table(rows=1, cols=2); st.style = "Table Grid"; st.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, t in enumerate(("Name of Student", "PRN")):
        st.rows[0].cells[i].text = ""; st.rows[0].cells[i].paragraphs[0].add_run(t).bold = True
    for s in cfg["students"]:
        c = st.add_row().cells; c[0].text, c[1].text = s["name"], s["prn"]
    D.d.add_paragraph()
    D.center(f"Div: {cfg['division']}", 10, True); D.center(f"Academic Year: {cfg['academic_year']}", 10, True)
    if not results:
        D.center("DRAFT TEMPLATE — result sections are pending a pipeline run on the real dataset.", 10, True, "C00000")
    D.pagebreak()

    # TOC placeholder (filled from headings collected in a dry pass)
    toc_anchor = D.d.add_paragraph(); toc_anchor.add_run("TABLE OF CONTENTS").bold = True
    toc_tbl_par = D.d.add_paragraph()
    D.pagebreak()

    # ---- 1
    H1("1. Introduction"); H2("1.1 Background")
    D.p("Air pollution is one of India's most significant environmental health risks. The Central Pollution Control Board (CPCB) operates a "
        "Continuous Ambient Air Quality Monitoring (CAAQM) network that reports pollutant concentrations hourly. Combined over many years and stations, "
        "these readings form a dataset large enough to need distributed processing for cleaning, aggregation and anomaly screening.")
    D.p("This Big Data Analytics mini project uses Apache Spark (PySpark) to ingest the India Air Quality Database, clean it, and compute city, state, "
        "seasonal, hourly, hotspot, anomaly and AQI summaries stored as Parquet. A Streamlit dashboard serves the saved results.")
    H2("1.2 Problem Statement")
    D.p("Develop a reproducible, laptop-runnable pipeline that analyses hourly air-quality readings across India, identifies pollution hotspots and "
        "abnormal pollution levels, compares cities and states, characterises monthly, seasonal and diurnal patterns, and presents the results in an interactive dashboard.")
    H2("1.3 Objectives")
    D.bullets(["Ingest and validate real CPCB-derived hourly data with Apache Spark using an explicit schema contract.",
               "Clean, standardise units, handle invalid values and deduplicate readings; report data-quality counts.",
               "Analyse PM2.5, PM10, NO2, SO2, CO, ozone and the other available pollutants.",
               "Rank cities/states and flag hotspot stations and abnormal pollution days using documented statistical rules.",
               "Compute the CPCB National AQI only where its validity rules are met.",
               "Provide an interactive Streamlit/Plotly dashboard with CSV downloads, plus logging and automated tests."])
    H2("1.4 Scope")
    if S:
        D.p(f"The analysed sample contains {S['rows_clean_unique']:,} cleaned hourly readings from {S['stations']} stations in {S['cities']} cities and "
            f"{S['states']} states between {S['first_timestamp_ist'][:10]} and {S['last_timestamp_ist'][:10]} (IST). The full database is much larger; "
            "the scope here is whatever date range is downloaded and analysed locally. Apache Cassandra was not implemented (Section 6).")
    else:
        D.p("The scope is the date range downloaded and analysed locally; coverage figures are filled in after the pipeline run. Apache Cassandra was not implemented (Section 6).")

    # ---- 2
    H1("2. Dataset and Preprocessing"); H2("2.1 Dataset Source and Collection")
    D.p("The data come from the India Air Quality Database published by the XKDR Forum (https://airquality.xkdr.org/). According to the project's "
        "site (checked 20 October 2026), it provides hourly readings from every station in CPCB's CAAQM network plus five US Embassy monitors: about 196.5 million "
        "readings from 558 stations and 15 pollutants, covering January 2009 to March 2026. Data are distributed through a query API and as one sorted Parquet file per month "
        "(under half a gigabyte in total).")
    D.p("Licence: Creative Commons Attribution 4.0 (CC BY 4.0), with credit to the India Air Quality Database (XKDR Forum) and the original sources "
        "(CPCB CAAQM; US Department of State via AirNow). The data are provided without warranty and are published as received from the networks, which label readings preliminary and not fully validated.")
    D.p("Download strategy for a student laptop: scripts/download_data.py fetches only the requested year(s)/months. A public demo key published on the site can read the 2024 files; "
        "a free personal key (https://airquality.xkdr.org/signup) has no row caps and reads all years. One year of Parquet is a small download relative to the whole archive.")
    H2("2.2 Important Attributes")
    D.table("Source schema (per the project's API documentation) and use in this project.", pd.DataFrame([
        ("station_id", "text", "Station key (CPCB site number or embassy id); identifies series and station counts."),
        ("station_name, state_name, city_name", "text", "Geography for city/state comparison and hotspot labels."),
        ("parameter_name", "text", "Pollutant (PM2.5, PM10, NO, NO2, NOx, NH3, SO2, CO, Ozone, Benzene, Toluene, Xylene, O_Xylene, MP-Xylene, Eth-Benzene)."),
        ("unit", "text", "Harmonised units (µg/m³, mg/m³, ppb); standardised in cleaning."),
        ("collected_at", "timestamp", "Naive timestamp in Indian Standard Time; CPCB readings sit on the hour."),
        ("value", "double", "Measured concentration, as received (no gap filling, outlier removal or calibration)."),
        ("source", "text", "cpcb_caaqm or us_embassy."),
        ("year, month", "int", "Hive partition columns of the bulk files; used for partition pruning."),
        ("latitude, longitude (stations.parquet)", "double", "Station coordinates for the hotspot map; some decommissioned stations have none.")],
        columns=["Attribute", "Type", "Analytical purpose"]), widths=[4.2, 1.8, 10])
    H2("2.3 Observed Dataset Coverage")
    if S:
        D.p(f"Spark read {S['rows_read']:,} rows. After cleaning, {S['rows_clean_unique']:,} unique station-pollutant-hour readings remain from {S['stations']} stations, "
            f"{S['cities']} cities and {S['states']} states, spanning {S['first_timestamp_ist']} to {S['last_timestamp_ist']} IST. "
            f"The year filter applied was {S['years_filter'] or 'none'}.")
    else:
        pend()
    H2("2.4 Schema-Aware Loading")
    D.p("read_raw() reads the monthly Parquet files with partition discovery, then enforces a schema contract: required columns (station_id, parameter_name, unit, collected_at, value) "
        "must exist or a SchemaError is raised; all columns are cast to canonical types (strings trimmed, value as double, collected_at as timestamp); optional columns that are absent become nulls. "
        "Parquet carries its own physical schema, so no inference over the data is needed. The Spark session timezone is Asia/Kolkata because the source timestamps are naive IST, so hour-of-day and date derivations are IST.")
    D.figure(figdir / "pipeline.png", "Spark cleaning and aggregation stages (data-quality stages in orange).")
    H2("2.5 Cleaning and Standardisation")
    D.bullets(["Rejected: missing station/parameter/timestamp (null_key_field); parameters outside the known list (unknown_parameter); null, NaN or negative values (invalid_value).",
               "Units normalised to µg/m³ (CO in mg/m³); µg↔mg conversions applied; ppb/ppm rows are quarantined (unit_not_convertible) rather than converted with assumed temperature/pressure.",
               "Rejected when above a physical plausibility bound per pollutant (above_physical_bound); bounds are configurable assumptions in src/aq/config.py, not standards.",
               "Deduplicated on (station_id, parameter_name, collected_at); conflicting duplicates are averaged and counted.",
               "Daily means require ≥ 16 hourly readings (CPCB rule); 8-hour windows for CO and ozone require ≥ 6 readings.",
               "No gap filling or outlier removal is applied; extreme-but-plausible values are kept and handled by the anomaly module."])
    H2("2.6 Cleaning Counts and Provenance")
    if S:
        rows = [("Rows read", S["rows_read"])] + [(f"Rejected: {k}", v) for k, v in S["rows_rejected_by_reason"].items()] + [
            ("Rows kept before deduplication", S["rows_kept_before_dedup"]), ("Duplicate keys removed", S["duplicates_removed"]),
            ("Conflicting duplicate keys (values differ)", S["conflicting_duplicate_keys"]), ("Clean unique hourly readings", S["rows_clean_unique"]),
            ("Stations / cities / states", f"{S['stations']} / {S['cities']} / {S['states']}"), ("Valid station-days (≥16 h)", S["station_days_valid"]),
            ("Station-days with valid CPCB AQI", S["station_days_with_valid_aqi"])]
        D.table("Data-quality counts from summary.json.", pd.DataFrame(rows, columns=["Check", "Observed value"]), widths=[9, 5])
        um = pd.DataFrame(S["unit_mix"]).rename(columns={"parameter_name": "Parameter", "unit_code": "Unit code", "n": "Rows"})
        D.table("Unit codes observed per parameter (raw rows).", um, widths=[4, 4, 4])
    else:
        pend()
    H2("2.7 Collection Limitations")
    D.p("Monitoring stations are not uniformly distributed and are concentrated in cities; readings are preliminary and not fully validated by the networks; "
        "station availability varies over time; and aggregates are not population-weighted. Findings describe monitored locations in the analysed period only.")

    # ---- 3
    H1("3. System Architecture and Technologies"); H2("3.1 Overall Architecture")
    D.figure(figdir / "architecture.png", "Architecture: download → Spark validation, cleaning and aggregation → Parquet/CSV → dashboard and report.", 13.5)
    H2("3.2 Layer Responsibilities")
    D.bullets(["Acquisition: scripts/download_data.py (year/month filter, retries, resume by size).",
               "Processing: src/aq/spark_pipeline.py — all cleaning and aggregation run as Spark DataFrame jobs.",
               "Storage: Parquet (clean hourly data partitioned by year/month; one Parquet dataset per aggregate) plus CSV exports and summary.json.",
               "Insights: src/aq/insights.py derives text only from computed tables.",
               "Presentation: app.py (Streamlit/Plotly) reads the aggregates only; Spark is not run per interaction."])
    H2("3.3 Technology Roles")
    D.table("Technology roles.", pd.DataFrame([
        ("Apache Spark / PySpark 4", "Schema enforcement, cleaning, window functions (8-hour averages, ranking), group-by aggregations, percentiles."),
        ("Parquet (pyarrow)", "Columnar input/output; partition pruning by year."),
        ("pandas", "Reading small aggregate tables for the dashboard, report and insights."),
        ("Streamlit + Plotly", "Interactive dashboard, filters, maps, downloads."),
        ("Matplotlib + python-docx", "Static report charts and the DOCX report."),
        ("pytest", "Automated unit, integration and dashboard smoke tests.")], columns=["Technology", "Role"]), widths=[4.5, 11.5])
    H2("3.4 Execution and Storage Boundaries")
    D.p("Spark runs in local mode (default local[2], 8 shuffle partitions, 2 GB driver memory; adjustable by command-line flags). Local threads demonstrate Spark's APIs and lazy execution "
        "on a laptop; they do not demonstrate cluster scalability, and no throughput benchmark is claimed. Aggregate tables are small and converted to pandas only after Spark has reduced them.")
    H2("3.5 Result Data Contract")
    D.p("Each aggregate is stored as results/parquet/<name>/ and results/csv/<name>.csv: city_summary, state_summary, monthly_city, monthly_national, seasonal_city, hourly_pattern, "
        "hotspots, anomaly_days, anomaly_summary, aqi_station_daily (Parquet only), aqi_city_monthly, aqi_category_city, aqi_dominant_city, data_coverage. summary.json holds run metadata and data-quality counts.")

    # ---- 4
    H1("4. Batch Analysis Using Apache Spark"); H2("4.1 Transformations and Actions")
    D.p("Transformations (select, filter, withColumn, groupBy, window functions, join) build a lazy plan; actions (write, count, first, collect, toPandas on aggregates) trigger execution. "
        "Key operations: groupBy().agg() for daily means with hour counts; percentile_approx for medians and 95th percentiles; Window.rangeBetween(-7h, 0) for 8-hour rolling means of CO and ozone; "
        "percent_rank and row_number windows for hotspot ranking; max_by for the dominant AQI pollutant.")
    H2("4.2 Counting Rules and Denominators")
    D.p("City/state statistics are means of valid station-day means, so every station-day has equal weight; they are not population-weighted and not hourly-reading weighted. "
        f"Cities and states are ranked only with at least the minimum station-days set in the dashboard (default 30); hotspot ranking needs at least {C.MIN_STATION_DAYS} valid days per station. "
        "Exceedance percentages use NAAQS 24-hour limits (PM2.5 60, PM10 100, NO2 80, SO2 80, NH3 400 µg/m³) applied to valid daily means only.")
    H2("4.3 Output Preparation")
    D.p("Every aggregate is written once as Parquet by Spark and exported to CSV for downloads; summary.json and pipeline.log record the run.")
    H2("4.4 Pollutant Coverage")
    if T:
        cov = T["data_coverage"].copy(); cov["first_date"] = cov.first_date.astype(str); cov["last_date"] = cov.last_date.astype(str)
        D.table("Valid station-days and station counts per pollutant.", cov.sort_values("n_station_days", ascending=False).rename(columns={
            "parameter_name": "Parameter", "n_stations": "Stations", "n_station_days": "Station-days", "first_date": "First", "last_date": "Last"}), widths=[3, 2.5, 3, 3, 3])
        D.p("Pollutants with few stations or station-days should not be compared with widely monitored ones; the parameter set available depends on station instrumentation.", "Interpretation.")
    else:
        pend()
    H2("4.5 City Comparison")
    if T and "cities" in ch:
        D.figure(ch["cities"], "City ranking by mean daily PM2.5 (cities with ≥ 30 valid station-days).")
        cs = T["city_summary"]; cs = cs[(cs.parameter_name == "PM2.5") & (cs.n_station_days >= 30)].sort_values("mean_daily", ascending=False).head(10)
        D.table("Top ten cities by mean daily PM2.5.", cs[["state_name", "city_name", "mean_daily", "median_daily", "p95_daily", "n_stations", "n_station_days", "pct_days_above_naaqs"]].rename(columns={
            "state_name": "State", "city_name": "City", "mean_daily": "Mean", "median_daily": "Median", "p95_daily": "P95", "n_stations": "Stations", "n_station_days": "Station-days", "pct_days_above_naaqs": "% > NAAQS"}), size=8)
        D.p(pick(ins, "Most polluted", "PM2.5:", "PM10:"), "Interpretation.")
    else:
        pend()
    H2("4.6 State Comparison")
    if T and "states" in ch:
        D.figure(ch["states"], "State ranking by mean daily PM2.5."); D.p(pick(ins, "Highest-"), "Interpretation.")
    else:
        pend()
    H2("4.7 Monthly Trend")
    if T and "monthly" in ch:
        D.figure(ch["monthly"], "Monthly mean of valid daily PM2.5 across stations.")
        mn = T["monthly_national"]; mn = mn[mn.parameter_name == "PM2.5"].sort_values("mean_daily")
        D.p(f"The monthly mean is lowest in {mn.iloc[0].ym} ({mn.iloc[0].mean_daily:.1f} µg/m³) and highest in {mn.iloc[-1].ym} ({mn.iloc[-1].mean_daily:.1f} µg/m³) across {len(mn)} month(s). "
            "The station set can change from month to month, so month-to-month changes combine real change with changing station coverage.", "Interpretation.")
    else:
        pend()
    H2("4.8 Seasonal Pattern")
    if T and "seasonal" in ch:
        D.figure(ch["seasonal"], "Seasonal mean daily PM2.5 (Winter Dec–Feb, Summer Mar–May, Monsoon Jun–Sep, Post-monsoon Oct–Nov).")
        D.p(pick(ins, "Seasonality") + " Seasons missing from the chart were not covered by the downloaded period.", "Interpretation.")
    else:
        pend()
    H2("4.9 Hourly Pattern")
    if T and "hourly" in ch:
        D.figure(ch["hourly"], "Average PM2.5 by hour of day (IST), reading-weighted across selected cities.")
        D.p(pick(ins, "Daily cycle"), "Interpretation.")
    else:
        pend()
    H2("4.10 Pollution Hotspots")
    if T:
        hs = T["hotspots"]; hs = hs[(hs.parameter_name == "PM2.5")].sort_values("rank").head(10)
        D.p(f"A station is a hotspot for a pollutant when its mean valid daily value is in the top {int((1 - C.HOTSPOT_PERCENTILE) * 100)}% of eligible stations (≥ {C.MIN_STATION_DAYS} valid days). "
            "This is a statistical definition for screening, not an official designation.")
        cols = [c for c in ["rank", "station_name", "city_name", "state_name", "mean_daily", "p95_daily", "n_days", "pct_days_above_naaqs", "is_hotspot"] if c in hs]
        D.table("Ten highest-ranked PM2.5 stations.", hs[cols].rename(columns={"station_name": "Station", "city_name": "City", "state_name": "State", "mean_daily": "Mean", "p95_daily": "P95", "n_days": "Days", "pct_days_above_naaqs": "% > NAAQS", "is_hotspot": "Hotspot"}), size=8)
        if "hotspot_map" in ch:
            D.figure(ch["hotspot_map"], "Station locations coloured by mean daily PM2.5; hotspots circled.", 9)
        D.p(pick(ins, "Hotspots", "States hosting"), "Interpretation.")
    else:
        pend()
    H2("4.11 Abnormal Pollution Levels")
    if T:
        D.p(f"For each station and pollutant (≥ {C.ANOMALY_MIN_DAYS} valid days) the modified z-score 0.6745·(x − median)/MAD is computed on daily means; days with z > {C.ANOMALY_Z} (Iglewicz–Hoaglin cut-off) are flagged as positive anomalies.")
        if "anomalies" in ch:
            D.figure(ch["anomalies"], "Cities with the most PM2.5 spike days.")
        ad = T["anomaly_days"]; ad = ad[ad.parameter_name == "PM2.5"].head(8)
        if len(ad):
            ad = ad.assign(date=ad.date.astype(str))
            D.table("Largest PM2.5 spike days.", ad[["station_name", "city_name", "date", "daily_mean", "station_median", "robust_z"]].rename(columns={"station_name": "Station", "city_name": "City", "date": "Date", "daily_mean": "Daily mean", "station_median": "Station median", "robust_z": "Robust z"}), size=8)
        D.p(pick(ins, "Anomalies") + " Flagged days may reflect real events (burning, dust, festivals) or sensor faults; they are screening results.", "Interpretation.")
    else:
        pend()
    H2("4.12 Air Quality Index")
    D.p("The database has no AQI column, so the CPCB National AQI is computed: sub-index linear within each band, 24-hour means for PM2.5, PM10, NO2, SO2, NH3 and 8-hour means for CO and ozone; "
        "overall AQI = highest sub-index; a station-day needs ≥ 16 hours per pollutant, ≥ 3 pollutants including PM2.5 or PM10. Band tables are implemented from the CPCB NAQI summary; "
        "CPCB does not publish the upper bound of the last band, so it is extended by the previous band's width and capped at 500 (assumption).")
    if T and "aqi" in ch:
        D.figure(ch["aqi"], "Computed CPCB AQI categories (station-days).")
        D.p(pick(ins, "Computed CPCB", "Dominant pollutant"), "Interpretation.")
    elif T:
        D.p("No station-day satisfied the AQI validity rules in this run.")
    else:
        pend()
    H2("4.13 Insights and Recommendations")
    if ins:
        D.p("Generated from computed tables by src/aq/insights.py:")
        D.bullets(ins["recommendations"] or ["No rule-based recommendation was triggered by the computed results."])
    else:
        pend()

    # ---- 5
    H1("5. Dashboard and Visualization"); H2("5.1 Streamlit Application")
    D.p("app.py reads the Parquet aggregates and presents eight tabs: Cities & states, Trends, Hourly, Hotspots, Anomalies, AQI, Insights and Data quality. Plotly provides interactive charts; every table has a CSV download button.")
    H2("5.2 Metric Cards")
    D.p("Cards show stations, cities/states, clean hourly readings, mean daily value of the selected pollutant, most polluted city, percentage of days above NAAQS, hotspot stations, anomalous station-days and valid-AQI station-days.")
    H2("5.3 Filters")
    D.p("Sidebar filters: pollutant, state, city, month range, season, top-N and minimum station-days. Aggregates are independent marginal tables, so a filter applies only to tables that carry that dimension (for example, the season filter applies to the seasonal view).")
    H2("5.4 Dashboard Screenshots")
    shots = ROOT / "docs" / "screenshots"
    for f, cap in SHOTS.items():
        if (shots / f).exists():
            D.figure(shots / f, cap, 15)
        else:
            D.p(f"[Screenshot pending: capture the dashboard tab and save as docs/screenshots/{f}] — {cap}", italic=True)
    H2("5.5 Downloads")
    D.p("Each table can be downloaded as CSV; the Data-quality tab also offers summary.json. The pipeline additionally writes all aggregates to results/csv/.")
    H2("5.6 Visualization Provenance")
    D.p("Figures in Section 4 are rendered by report/build_report.py from results/parquet; dashboard screenshots are taken from the running application. Diagrams are drawn by report/diagrams.py.")

    # ---- 6
    H1("6. Storage and Integration"); H2("6.1 Parquet Storage Layout")
    D.table("Storage layout.", pd.DataFrame([
        ("data/raw/measurements/year=/month=", "Downloaded source Parquet (read-only)."),
        ("results/parquet/clean_hourly/year=/month=", "Cleaned, deduplicated hourly readings in canonical units."),
        ("results/parquet/<aggregate>/", "One Parquet dataset per aggregate table."),
        ("results/csv/<aggregate>.csv", "CSV exports for download and the report."),
        ("results/summary.json, pipeline.log, insights.md/json", "Run metadata, logs, generated insights.")], columns=["Path", "Content"]), widths=[7, 9])
    H2("6.2 Why Parquet")
    D.p("Parquet is columnar, compressed and natively read/written by Spark and pandas, which keeps the project runnable on a laptop without a database server.")
    H2("6.3 Apache Cassandra (Not Implemented)")
    D.p("Cassandra was optional and is not part of this implementation. The workload is batch analytics over a static historical dataset, for which Parquet is sufficient. "
        "Cassandra would be appropriate for serving low-latency lookups (for example, latest readings per station) in a future streaming extension; this is future scope only.")
    H2("6.4 Reliability Considerations")
    D.p("Aggregate tables are overwritten on each run (not versioned), so an interrupted run can leave a mix of old and new tables; re-run the pipeline to restore consistency. The pipeline exits with a non-zero code and a logged message for missing input, schema errors or no valid rows.")

    # ---- 7
    H1("7. Local Deployment and Execution"); H2("7.1 Development Environment")
    D.p("Python 3.10+, Java 17 or later (required by Spark), 8 GB RAM recommended. Dependencies are in requirements.txt (PySpark 4.0.1, pandas, pyarrow, Streamlit, Plotly, python-docx, matplotlib, pytest).")
    H2("7.2 Prepare the Data")
    D.code("python -m venv .venv && source .venv/bin/activate        # Windows: .venv\\Scripts\\activate\npip install -r requirements.txt\npython scripts/download_data.py --years 2024")
    H2("7.3 Run the Spark Pipeline")
    D.code("python -m src.aq.spark_pipeline --years 2024 --master local[4] --driver-memory 4g\npython -m src.aq.insights results")
    H2("7.4 Start the Dashboard")
    D.code("streamlit run app.py")
    H2("7.5 Build the Report and Run Tests")
    D.code("pytest --junitxml=docs/pytest_report.xml\npython report/build_report.py --results results")

    # ---- 8
    H1("8. Testing and Validation"); H2("8.1 Validation Approach")
    D.p("Automated pytest tests verify cleaning rules, unit conversion, deduplication, schema errors, AQI maths (including Spark/Python parity), end-to-end aggregates against hand-computed expectations, "
        "hotspot and anomaly detection, insights generation and a Streamlit smoke test (AppTest). The end-to-end tests run on a small deterministic synthetic fixture created in tests/conftest.py, "
        "chosen so every expected value can be computed by hand; it is used for testing only and never for reported findings.")
    H2("8.2 Timestamp Handling")
    D.p("During development, a test exposed a timezone defect: collecting a Spark timestamp into Python shifted it by the machine's UTC offset (5 h 30 min relative to IST). "
        "Timestamp bounds in summary.json are therefore formatted inside Spark in the IST session timezone, and the regression test asserts the expected bound.")
    H2("8.3 Test Results")
    if jr:
        n = len(jr); bad = sum(r[2] != "passed" for r in jr)
        D.p(f"The suite was executed on {jattr.get('timestamp', 'n/a')[:19]} (development sandbox, Linux): {n} tests, {bad} failed, in {float(jattr.get('time', 0)):.0f} s.")
        D.table("Automated test results (from docs/pytest_report.xml).", pd.DataFrame(jr, columns=["Module", "Test", "Result"]), widths=[4.5, 9, 2], size=8)
        D.p("These results come from the synthetic test fixture. They validate code logic, not the real dataset; re-run pytest on your machine.")
    else:
        pend()
    H2("8.4 Data-Quality Checks on Real Data")
    D.p("Every pipeline run records rejected-row counts, duplicates and unit mix (Section 2.6). " + ("" if S else "Real-data values are pending a pipeline run."))
    H2("8.5 Limitations of Validation")
    D.bullets(["Download from airquality.xkdr.org was not possible in the development sandbox (network policy); the downloader was written from the published API documentation and unit-tested only for file selection.",
               "Real-data runtime, memory use and scalability were not measured there; local[2] results on a small fixture do not predict full-year performance.",
               "AQI band tables and NAAQS values were taken from public summaries of CPCB documents and should be confirmed against the official notifications for formal use."])

    # ---- 9
    H1("9. Conclusion and Future Scope"); H2("9.1 Project Summary")
    if ins:
        D.p("Key computed findings:"); D.bullets(ins["findings"])
        D.p("Limitations of these findings:"); D.bullets(ins["limitations"])
    else:
        D.p("The project delivers a PySpark pipeline, Parquet storage, a Streamlit dashboard, tests and a report generator. Findings are produced after running the pipeline on the real dataset.")
    H2("9.2 Future Scope")
    D.bullets(["Run on multi-year data and add a cluster/Docker-based Spark profile with measured scalability.", "Add Apache Cassandra or another serving store for near-real-time station lookups.",
               "Population- or area-weighted aggregation and meteorological covariates.", "Forecasting models and source apportionment for hotspot stations.",
               "Versioned result snapshots and automated monthly refresh."])
    H1("10. References")
    refs = ["XKDR Forum, India Air Quality Database, https://airquality.xkdr.org (CC BY 4.0); source: CPCB CAAQM network and US Department of State via AirNow.",
            "xKDR/Air-Quality-Database repository and docs/PUBLIC_API.md, https://github.com/xKDR/Air-Quality-Database.",
            "Central Pollution Control Board, National Air Quality Index (summary: https://ies.gov.in/arthapedia/concept/national-air-quality-index).",
            "National Ambient Air Quality Standards 2009 (summary: https://ies.gov.in/arthapedia/concept/ambient-air-quality-standards-india).",
            "Iglewicz, B. and Hoaglin, D., How to Detect and Handle Outliers, ASQC Quality Press, 1993 (modified z-score).",
            "Apache Spark / PySpark documentation; Streamlit and Plotly documentation."]
    for i, r in enumerate(refs, 1):
        D.p(f"{i}. {r}")

    # ---- fill TOC
    tbl = D.d.add_table(rows=0, cols=2)
    for lvl, text in heads:
        c = tbl.add_row().cells
        c[0].text = ""; run = c[0].paragraphs[0].add_run(("    " if lvl == 2 else "") + text); run.bold = lvl == 1; run.font.size = Pt(10)
        c[1].text = ""; c[1].paragraphs[0].add_run(str(toc_pages.get(text, "–"))).font.size = Pt(10); c[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        c[0].width, c[1].width = Cm(14), Cm(2)
        for cell in c:
            cell.paragraphs[0].paragraph_format.space_after = Pt(0); cell.paragraphs[0].paragraph_format.space_before = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
    toc_tbl_par._p.addnext(tbl._tbl)
    D.d.save(out_docx)
    return heads


def to_pdf(docx: Path) -> Path:
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
    return docx.with_suffix(".pdf")


def page_map(pdf: Path, heads) -> dict:
    from pypdf import PdfReader
    pages = [re.sub(r"\s+", " ", (p.extract_text() or "")) for p in PdfReader(str(pdf)).pages]
    # body starts at the page holding the first Introduction paragraph; everything before is cover + TOC
    skip_pages = next((i for i, t in enumerate(pages) if "Air pollution is one of India" in t), 0)
    out = {}
    for _, text in heads:
        norm = re.sub(r"\s+", " ", text)
        for i, t in enumerate(pages):
            if i >= skip_pages and norm in t:
                out[text] = i + 1
                break
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--results", default=str(C.DEFAULT_RESULTS)); ap.add_argument("--template", action="store_true")
    ap.add_argument("--config", default=str(ROOT / "report" / "report_config.json")); ap.add_argument("--out", default=str(ROOT / "report" / "output"))
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args(argv)
    res = None if a.template else Path(a.results)
    if res and not (res / "summary.json").exists():
        print(f"No pipeline results in {res}. Run the pipeline first (README), or use --template.", file=sys.stderr)
        return 2
    cfg = json.loads(Path(a.config).read_text(encoding="utf-8")); out = Path(a.out); out.mkdir(parents=True, exist_ok=True)
    name = "BDA_Mini_Project_Report" + ("_TEMPLATE" if a.template else "")
    docx = out / f"{name}.docx"
    heads = build(cfg, res, {}, docx, out / "figures")
    if a.no_pdf:
        return 0
    pdf = to_pdf(docx)
    build(cfg, res, page_map(pdf, heads), docx, out / "figures")     # second pass with real page numbers
    pdf = to_pdf(docx)
    print("wrote", docx, "and", pdf)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
